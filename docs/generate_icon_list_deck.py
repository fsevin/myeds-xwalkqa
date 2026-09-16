from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "icon-list-implementation.pptx"

RED = RGBColor(235, 16, 0)
BLACK = RGBColor(25, 25, 25)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(46, 46, 46)
MUTED = RGBColor(108, 108, 108)
PANEL = RGBColor(246, 246, 246)
LINE = RGBColor(220, 220, 220)
CODE = RGBColor(35, 35, 35)
GREEN = RGBColor(28, 118, 72)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font="Arial", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_text(slide, runs, x, y, w, h, size=16, color=INK, font="Arial"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    for text, bold, run_color in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = run_color or color
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_rule(slide, x, y, w, color=LINE, thickness=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(thickness / 72))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_footer(slide, number):
    add_rule(slide, 0.6, 7.08, 12.1, RED, 2)
    add_text(slide, "ADOBE  |  EDGE DELIVERY SERVICES", 0.6, 7.18, 5, 0.2, 8, MUTED, True)
    add_text(slide, f"ICON-LIST IMPLEMENTATION  •  {number:02d}", 9.1, 7.18, 3.6, 0.2, 8, MUTED, True, align=PP_ALIGN.RIGHT)


def title(slide, eyebrow, heading, subheading=None):
    add_text(slide, eyebrow.upper(), 0.65, 0.45, 6, 0.25, 10, RED, True)
    add_text(slide, heading, 0.65, 0.82, 11.8, 0.7, 30, BLACK, True)
    if subheading:
        add_text(slide, subheading, 0.67, 1.62, 11, 0.4, 14, MUTED)


def code_panel(slide, code, x, y, w, h, label=None, size=11):
    add_rect(slide, x, y, w, h, CODE, CODE, True)
    if label:
        add_text(slide, label.upper(), x + 0.22, y + 0.16, w - 0.44, 0.2, 8, RGBColor(180, 180, 180), True)
        code_y = y + 0.46
    else:
        code_y = y + 0.2
    add_text(slide, code, x + 0.22, code_y, w - 0.44, h - (code_y - y) - 0.18, size, WHITE, font="Courier New")


def bullet_list(slide, items, x, y, w, size=17, gap=0.45, color=INK):
    for index, item in enumerate(items):
        yy = y + index * gap
        add_text(slide, "•", x, yy, 0.22, 0.25, size, RED, True)
        add_text(slide, item, x + 0.3, yy, w - 0.3, 0.35, size, color)


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, BLACK)
    add_rect(slide, 0, 0, 0.22, 7.5, RED)
    add_text(slide, "ADOBE", 0.75, 0.62, 2.1, 0.4, 22, WHITE, True)
    add_text(slide, "Edge Delivery Services", 0.77, 1.67, 5.4, 0.35, 14, RGBColor(190, 190, 190), True)
    add_text(slide, "Icon List\nImplementation", 0.75, 2.05, 7.2, 1.65, 42, WHITE, True)
    add_text(slide, "How authored rows become responsive, optimized list items", 0.78, 4.02, 6.6, 0.5, 18, RGBColor(220, 220, 220))
    add_rect(slide, 8.4, 1.4, 3.9, 4.55, RGBColor(49, 49, 49), RGBColor(85, 85, 85), True)
    for idx, (icon, heading, body) in enumerate([
        ("01", "Author", "Icon, alt text, heading, description"),
        ("02", "Decorate", "Rows become semantic <li> elements"),
        ("03", "Optimize", "Images load through AEM transforms"),
        ("04", "Style", "Flex layout adapts across viewports"),
    ]):
        yy = 1.72 + idx * 0.92
        add_text(slide, icon, 8.72, yy, 0.35, 0.25, 11, RED, True, align=PP_ALIGN.CENTER)
        add_text(slide, heading, 9.28, yy - 0.02, 1.3, 0.25, 15, WHITE, True)
        add_text(slide, body, 9.28, yy + 0.28, 2.55, 0.35, 10, RGBColor(185, 185, 185))
    add_text(slide, "myeds-xwalkqa  •  September 2026", 0.78, 6.78, 4, 0.2, 9, RGBColor(150, 150, 150))

    # 2. Architecture at a glance
    slide = prs.slides.add_slide(blank)
    title(slide, "01  |  Architecture", "One block, four cooperating layers", "The icon-list implementation follows the standard EDS convention: authored data, block decoration, shared utilities, and CSS.")
    layers = [
        ("Authoring model", "_icon-list.json", "Defines the item fields and authoring filter.", RED),
        ("Block decorator", "icon-list.js", "Transforms authored DOM into semantic list markup.", BLACK),
        ("Shared runtime", "scripts/aem.js", "Provides image optimization and dynamic loading.", RGBColor(76, 76, 76)),
        ("Presentation", "icon-list.css", "Controls responsive layout, type, and spacing.", RGBColor(135, 135, 135)),
    ]
    for i, (heading, file, body, fill) in enumerate(layers):
        x = 0.8 + i * 3.05
        add_rect(slide, x, 2.45, 2.55, 2.25, fill, fill, True)
        add_text(slide, f"0{i + 1}", x + 0.22, 2.68, 0.5, 0.25, 11, WHITE, True)
        add_text(slide, heading, x + 0.22, 3.15, 2.1, 0.5, 18, WHITE, True)
        add_text(slide, file, x + 0.22, 3.83, 2.1, 0.3, 11, RGBColor(240, 240, 240), font="Courier New")
        add_text(slide, body, x + 0.22, 4.18, 2.05, 0.4, 10, RGBColor(240, 240, 240))
        if i < len(layers) - 1:
            add_text(slide, "→", x + 2.62, 3.35, 0.35, 0.35, 20, RED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "The contract is intentionally small: the runtime discovers the block from its class name and imports matching JS/CSS files.", 1.25, 5.55, 10.8, 0.55, 17, BLACK, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 2)

    # 3. Authoring model
    slide = prs.slides.add_slide(blank)
    title(slide, "02  |  Authoring", "The model defines one repeatable item", "The block exposes only the fields an author needs, while the filter prevents unrelated components from being inserted into the list.")
    code_panel(slide, '''{
  "id": "icon-list-item",
  "fields": [
    { "name": "image", "component": "reference" },
    { "name": "imageAlt", "component": "text" },
    { "name": "heading", "required": true },
    { "name": "text", "component": "richtext" }
  ]
}''', 0.75, 2.3, 5.55, 3.75, "blocks/icon-list/_icon-list.json", 12)
    add_text(slide, "Authoring contract", 7.0, 2.38, 4.7, 0.35, 21, BLACK, True)
    bullet_list(slide, [
        "Icon is a single referenced image.",
        "Alt text is modeled separately for accessibility.",
        "Heading is required and becomes the item title.",
        "Description supports rich text content.",
        "The icon-list filter allows only icon-list-item entries.",
    ], 7.02, 3.02, 5.2, 15, 0.54)
    add_rect(slide, 7.02, 5.86, 4.95, 0.48, PANEL, LINE, True)
    add_text(slide, "Source of truth: blocks/icon-list/_icon-list.json", 7.28, 6.01, 4.35, 0.18, 10, MUTED, font="Courier New")
    add_footer(slide, 3)

    # 4. Runtime transformation
    slide = prs.slides.add_slide(blank)
    title(slide, "03  |  Runtime", "Decoration converts rows into semantic markup", "The default decorator keeps the authored content, adds semantic structure, and preserves Universal Editor instrumentation.")
    code_panel(slide, '''const ul = document.createElement('ul');
[...block.children].forEach((row) => {
  const li = document.createElement('li');
  moveInstrumentation(row, li);
  while (row.firstElementChild) li.append(row.firstElementChild);
  ul.append(li);
});
block.replaceChildren(ul);''', 0.75, 2.2, 6.1, 3.4, "icon-list.js  •  core transformation", 12)
    add_text(slide, "Before", 7.45, 2.35, 1.2, 0.3, 12, MUTED, True)
    add_rect(slide, 7.45, 2.75, 4.75, 0.9, PANEL, LINE, True)
    add_text(slide, "div.icon-list  ›  div row  ›  div cells", 7.72, 3.05, 4.2, 0.25, 12, INK, font="Courier New")
    add_text(slide, "↓", 9.58, 3.83, 0.35, 0.35, 22, RED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "After", 7.45, 4.34, 1.2, 0.3, 12, MUTED, True)
    add_rect(slide, 7.45, 4.74, 4.75, 1.25, BLACK, BLACK, True)
    add_text(slide, "ul\n  └─ li\n      ├─ icon\n      ├─ heading\n      └─ description", 7.72, 4.92, 4.1, 0.92, 12, WHITE, font="Courier New")
    add_footer(slide, 4)

    # 5. Image optimization
    slide = prs.slides.add_slide(blank)
    title(slide, "04  |  Media", "Images use the shared optimization pipeline", "The block delegates image URL construction to createOptimizedPicture, keeping media behavior consistent across the site.")
    code_panel(slide, '''const optimizedPic = createOptimizedPicture(
  img.src,
  img.alt,
  false,
  [{ width: '100' }],
);
moveInstrumentation(img, optimizedPic.querySelector('img'));
img.closest('picture').replaceWith(optimizedPic);''', 0.75, 2.2, 6.45, 3.45, "icon-list.js  •  image path", 11)
    add_text(slide, "What this gives the block", 7.75, 2.36, 4.3, 0.35, 20, BLACK, True)
    bullet_list(slide, [
        "WebP source with a browser-compatible fallback.",
        "A compact 100px asset target for icon-sized media.",
        "Alt text carried through to the optimized image.",
        "Editor instrumentation retained after replacement.",
    ], 7.78, 3.02, 4.6, 15, 0.6)
    add_rect(slide, 7.78, 5.7, 4.35, 0.55, RGBColor(232, 247, 238), RGBColor(188, 225, 202), True)
    add_text(slide, "Shared utility  →  scripts/aem.js", 8.05, 5.88, 3.8, 0.18, 11, GREEN, True, font="Courier New")
    add_footer(slide, 5)

    # 6. CSS behavior
    slide = prs.slides.add_slide(blank)
    title(slide, "05  |  Presentation", "CSS creates a responsive icon grid", "The block owns its layout. Global variables supply the site-wide type scale and colors.")
    code_panel(slide, '''.icon-list ul {
  display: flex;
  flex-wrap: wrap;
  justify-content: center;
  gap: 40px 24px;
}

.icon-list li {
  flex: 1 1 240px;
  max-width: 360px;
  text-align: center;
}''', 0.75, 2.15, 5.35, 3.6, "icon-list.css  •  layout rules", 12)
    add_text(slide, "Desktop behavior", 6.85, 2.38, 2.2, 0.25, 12, MUTED, True)
    for row in range(2):
        for col in range(3):
            x = 6.9 + col * 1.75
            y = 2.82 + row * 0.85
            add_rect(slide, x, y, 1.38, 0.6, PANEL, LINE, True)
            add_rect(slide, x + 0.54, y + 0.08, 0.3, 0.3, RED, RED, True)
            add_text(slide, "heading", x + 0.15, y + 0.42, 1.1, 0.12, 7, MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "Fluid items wrap at 240px and cap at 360px.", 6.9, 4.65, 4.9, 0.3, 15, BLACK, True)
    add_text(slide, "Mobile behavior", 6.85, 5.22, 2.2, 0.25, 12, MUTED, True)
    add_rect(slide, 6.9, 5.62, 4.65, 0.55, PANEL, LINE, True)
    add_rect(slide, 7.1, 5.74, 4.25, 0.3, RED, RED, True)
    add_text(slide, "Items naturally collapse into a single readable column.", 6.9, 6.34, 5.0, 0.3, 13, INK)
    add_footer(slide, 6)

    # 7. Lifecycle
    slide = prs.slides.add_slide(blank)
    title(slide, "06  |  Lifecycle", "When the page loads, the block follows the EDS lifecycle", "The page bootstrapper decorates first, then loads block assets on demand so initial rendering stays lightweight.")
    steps = [
        ("HTML", "AEM delivers authored rows", "content"),
        ("Decorate", "aem.js adds .block and data-block-name", "runtime"),
        ("Import", "Matching JS/CSS are loaded", "assets"),
        ("Render", "decorate(block) replaces rows with <ul>", "DOM"),
        ("Loaded", "data-block-status becomes loaded", "state"),
    ]
    for i, (head, body, tag) in enumerate(steps):
        x = 0.8 + i * 2.47
        add_rect(slide, x, 2.55, 2.0, 2.1, BLACK if i in (1, 3) else PANEL, BLACK if i in (1, 3) else LINE, True)
        add_text(slide, f"0{i + 1}", x + 0.18, 2.78, 0.35, 0.2, 10, RED, True)
        add_text(slide, head, x + 0.18, 3.15, 1.6, 0.3, 17, WHITE if i in (1, 3) else BLACK, True)
        add_text(slide, body, x + 0.18, 3.64, 1.62, 0.55, 11, RGBColor(210, 210, 210) if i in (1, 3) else INK)
        add_text(slide, tag, x + 0.18, 4.3, 1.3, 0.16, 8, RGBColor(170, 170, 170) if i in (1, 3) else MUTED, True, font="Courier New")
        if i < len(steps) - 1:
            add_text(slide, "→", x + 2.08, 3.42, 0.3, 0.3, 18, RED, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 1.15, 5.55, 10.95, 0.62, RGBColor(255, 244, 242), RGBColor(255, 205, 198), True)
    add_text(slide, "Key property: the block owns its behavior, while aem.js owns discovery and lifecycle orchestration.", 1.45, 5.77, 10.35, 0.2, 14, RED, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)

    # 8. Implementation map
    slide = prs.slides.add_slide(blank)
    title(slide, "07  |  Implementation map", "Where to change the behavior", "A practical guide for extending the block without breaking the EDS contract.")
    rows = [
        ("Authoring fields", "blocks/icon-list/_icon-list.json", "Add or revise model fields and item filters."),
        ("DOM transformation", "blocks/icon-list/icon-list.js", "Change markup, classification, or image handling."),
        ("Visual design", "blocks/icon-list/icon-list.css", "Change layout, spacing, typography, or responsive behavior."),
        ("Shared image behavior", "scripts/aem.js", "Change only when the optimization contract should change site-wide."),
        ("Generated authoring output", "component-*.json", "Regenerate with npm run build:json after model changes."),
    ]
    y = 2.25
    for idx, (area, file, action) in enumerate(rows):
        fill = PANEL if idx % 2 == 0 else WHITE
        add_rect(slide, 0.75, y, 11.85, 0.67, fill, LINE)
        add_text(slide, area, 1.02, y + 0.17, 2.05, 0.2, 12, BLACK, True)
        add_text(slide, file, 3.28, y + 0.17, 4.0, 0.2, 11, RED, font="Courier New")
        add_text(slide, action, 7.55, y + 0.17, 4.65, 0.25, 11, INK)
        y += 0.75
    add_text(slide, "Recommended extension pattern", 0.8, 6.2, 3.0, 0.25, 14, BLACK, True)
    add_text(slide, "Model the content → decorate the authored DOM → reuse shared utilities → style locally → regenerate JSON.", 3.45, 6.2, 8.6, 0.25, 13, MUTED)
    add_footer(slide, 8)

    prs.save(OUTPUT)


if __name__ == "__main__":
    make_deck()
    print(OUTPUT)